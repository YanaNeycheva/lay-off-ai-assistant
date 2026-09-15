"""Export the person's live tracker.md to a point-in-time .xlsx snapshot.

Harness-independent replacement for the bundled ``xlsx`` skill. The markdown
``tracker.md`` stays the canonical source of truth (the orchestrator keeps it
live); this produces a clean spreadsheet copy the person can sort / filter /
share offline. It does NOT round-trip back — it is a snapshot, not a second
source of truth (see agent/templates/tracker.md "Hybrid format").

Each ``##`` section that contains a markdown table becomes one worksheet; the
"Инерция" momentum block (a bullet list) becomes a key/value summary sheet.

Usage:
    python scripts/tracker_to_xlsx.py <tracker.md> --out <tracker.xlsx>

Requires: openpyxl  (pip install openpyxl)
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# Excel forbids these in sheet names, and caps names at 31 chars.
_INVALID_SHEET_CHARS = re.compile(r"[\\/?*\[\]:]")
# Keep letters (incl. Cyrillic), digits, spaces and dashes; drop emoji/symbols.
_KEEP = re.compile(r"[^\w\s\-]", re.UNICODE)


def _clean_sheet_name(heading: str, used: set[str]) -> str:
    name = _KEEP.sub("", heading)
    name = _INVALID_SHEET_CHARS.sub("", name).strip()
    name = re.sub(r"\s+", " ", name) or "Раздел"
    name = name[:31]
    base, i = name, 2
    while name.lower() in used:
        suffix = f" {i}"
        name = (base[: 31 - len(suffix)] + suffix)
        i += 1
    used.add(name.lower())
    return name


def _split_row(line: str) -> list[str]:
    # "| a | b | c |" -> ["a", "b", "c"]
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def _is_separator(cells: list[str]) -> bool:
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c != "")


def parse_tracker(text: str) -> dict:
    """Parse tracker markdown into {sections: [(heading, table_rows)], momentum: {...}}."""
    lines = text.splitlines()
    sections: list[tuple[str, list[list[str]]]] = []
    momentum: dict[str, str] = {}

    heading = None
    in_momentum = False
    table: list[list[str]] = []

    def flush():
        nonlocal table
        if heading and table:
            sections.append((heading, table))
        table = []

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("## "):
            flush()
            heading = line[3:].strip()
            in_momentum = "инерц" in heading.lower()
            continue
        if line.startswith("#"):  # top title / other headings end a table
            flush()
            heading = None
            in_momentum = False
            continue
        if in_momentum and line.strip().startswith("- "):
            item = line.strip()[2:]
            item = item.replace("**", "")
            if ":" in item:
                key, val = item.split(":", 1)
                momentum[key.strip()] = val.strip()
            else:
                momentum[item.strip()] = ""
            continue
        if line.strip().startswith("|"):
            cells = _split_row(line)
            if _is_separator(cells):
                continue
            table.append(cells)
        elif table:
            # Blank / prose line ends the current table under this heading.
            flush()

    flush()
    return {"sections": sections, "momentum": momentum}


def _autosize(ws) -> None:
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        width = max((len(str(c.value)) for c in column_cells if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width + 2, 12), 60)


def build_workbook(parsed: dict) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    bold = Font(bold=True)

    momentum = parsed.get("momentum") or {}
    if momentum:
        ws = wb.create_sheet(_clean_sheet_name("Инерция", used))
        ws.append(["Показател", "Стойност"])
        for c in ws[1]:
            c.font = bold
        for key, val in momentum.items():
            ws.append([key, val])
        _autosize(ws)

    for heading, rows in parsed.get("sections") or []:
        ws = wb.create_sheet(_clean_sheet_name(heading, used))
        ws.append([heading])
        ws["A1"].font = bold
        if rows:
            ws.append(rows[0])  # header row
            for c in ws[2]:
                c.font = bold
            for row in rows[1:]:
                ws.append(row)
        _autosize(ws)

    if not wb.sheetnames:  # never leave an empty workbook
        wb.create_sheet("Tracker")
    return wb


def main(argv=None) -> int:
    try:  # keep Cyrillic status prints from crashing on a cp1252 console (Windows)
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    parser = argparse.ArgumentParser(description="Export tracker.md to an .xlsx snapshot")
    parser.add_argument("tracker_md", help="Path to the person's tracker.md")
    parser.add_argument("--out", required=True, help="Output .xlsx path")
    args = parser.parse_args(argv)

    text = Path(args.tracker_md).read_text(encoding="utf-8")
    parsed = parse_tracker(text)
    wb = build_workbook(parsed)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(f"Wrote {out} ({len(wb.sheetnames)} sheets)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
