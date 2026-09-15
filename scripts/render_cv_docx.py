"""Render a CV from a structured JSON file to an ATS-clean .docx.

Harness-independent replacement for the bundled ``docx`` skill: the agent
composes the CV *content* (per ``agent/templates/cv-structure.md``) and writes
it as JSON; this script renders it deterministically into an ATS-compliant
Word document. Determinism is the win — the ATS rules are enforced by code,
not re-derived by an LLM each time.

Contract (the JSON the agent produces). Everything except ``name`` is optional;
omitted sections are simply skipped:

    {
      "name": "Име Фамилия",
      "language": "bg" | "en",          # picks section-heading language; default "bg"
      "contact": {                        # rendered top, joined with " · "
        "location": "София",
        "phone": "+359 ...",
        "email": "...@...",
        "linkedin": "linkedin.com/in/..."
      },
      "summary": "3–4 line targeted summary.",
      "experience": [
        {
          "title": "Финансов анализатор",
          "company": "Acme",
          "location": "София",
          "start": "Март 2019",
          "end": "Юни 2024",           # or "настоящем" / "present"
          "bullets": ["Резултат с число ...", "..."]
        }
      ],
      "skills": ["Excel", "SQL"]          # list -> comma-joined; OR
                                          # {"Софтуер": ["Excel"], "Езици": ["EN"]}
      "education": [
        {"degree": "Магистър Финанси", "institution": "УНСС", "year": "2015"}
      ],
      "certifications": ["CFA Level I"]
    }

ATS rules honored (see .claude/skills/tailor-cv/ats-rules.md): single column,
no tables / text boxes / headers / footers / images, standard section headings,
contact block at the very top, en-dash (–) in date ranges. Whether bullets
start with an action verb and carry numbers is the agent's job, not the
renderer's — this only guarantees the *structure* is ATS-clean.

Usage:
    python scripts/render_cv_docx.py <cv.json> --out "<Name> CV.docx"

Requires: python-docx  (pip install python-docx)
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:  # pragma: no cover - exercised only without the dep
    sys.stderr.write(
        "python-docx is not installed. Install it with:\n"
        "    pip install python-docx\n"
    )
    raise

EN_DASH = "–"

# Standard ATS section headings, per ats-rules.md rule 2.
HEADINGS = {
    "bg": {
        "summary": "Резюме",
        "experience": "Опит",
        "skills": "Умения",
        "education": "Образование",
        "certifications": "Сертификати",
    },
    "en": {
        "summary": "Summary",
        "experience": "Work Experience",
        "skills": "Skills",
        "education": "Education",
        "certifications": "Certifications",
    },
}

BASE_FONT = "Calibri"


def _date_range(entry: dict) -> str:
    """Join start/end with an en-dash; tolerate a missing end."""
    start = str(entry.get("start", "")).strip()
    end = str(entry.get("end", "")).strip()
    if start and end:
        return f"{start} {EN_DASH} {end}"
    return start or end


def _heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11.5)


def _contact_line(contact: dict) -> str:
    order = ["location", "phone", "email", "linkedin"]
    parts = [str(contact.get(k, "")).strip() for k in order]
    return " · ".join(p for p in parts if p)  # middot separator


def build_document(cv: dict):
    lang = cv.get("language", "bg")
    headings = HEADINGS.get(lang, HEADINGS["bg"])

    doc = Document()

    # Clean, single-column defaults. No headers/footers are added (python-docx
    # leaves them empty), no tables are ever created — that keeps it ATS-clean.
    normal = doc.styles["Normal"]
    normal.font.name = BASE_FONT
    normal.font.size = Pt(10.5)

    # --- Name -------------------------------------------------------------
    name_p = doc.add_paragraph()
    name_p.paragraph_format.space_after = Pt(0)
    name_run = name_p.add_run(str(cv.get("name", "")).strip())
    name_run.bold = True
    name_run.font.size = Pt(18)

    # --- Contact ----------------------------------------------------------
    contact = cv.get("contact") or {}
    contact_text = _contact_line(contact)
    if contact_text:
        cp = doc.add_paragraph(contact_text)
        cp.paragraph_format.space_after = Pt(2)

    # --- Summary ----------------------------------------------------------
    summary = str(cv.get("summary", "")).strip()
    if summary:
        _heading(doc, headings["summary"])
        doc.add_paragraph(summary)

    # --- Experience -------------------------------------------------------
    experience = cv.get("experience") or []
    if experience:
        _heading(doc, headings["experience"])
        for job in experience:
            header_bits = [str(job.get("title", "")).strip(), str(job.get("company", "")).strip()]
            line = f"{header_bits[0]} {EN_DASH} {header_bits[1]}".strip(f" {EN_DASH}")
            meta = " · ".join(
                b for b in [str(job.get("location", "")).strip(), _date_range(job)] if b
            )
            if meta:
                line = f"{line} · {meta}" if line else meta
            jp = doc.add_paragraph()
            jp.paragraph_format.space_before = Pt(6)
            jp.paragraph_format.space_after = Pt(0)
            jp.add_run(line).bold = True
            for bullet in job.get("bullets") or []:
                doc.add_paragraph(str(bullet).strip(), style="List Bullet")

    # --- Skills -----------------------------------------------------------
    skills = cv.get("skills")
    if skills:
        _heading(doc, headings["skills"])
        if isinstance(skills, dict):
            for category, items in skills.items():
                joined = ", ".join(str(i).strip() for i in items)
                p = doc.add_paragraph()
                p.add_run(f"{category}: ").bold = True
                p.add_run(joined)
        else:
            doc.add_paragraph(", ".join(str(i).strip() for i in skills))

    # --- Education --------------------------------------------------------
    education = cv.get("education") or []
    if education:
        _heading(doc, headings["education"])
        for ed in education:
            bits = [str(ed.get("degree", "")).strip(), str(ed.get("institution", "")).strip()]
            line = f"{bits[0]} {EN_DASH} {bits[1]}".strip(f" {EN_DASH}")
            year = str(ed.get("year", "")).strip()
            if year:
                line = f"{line} · {year}" if line else year
            doc.add_paragraph(line)

    # --- Certifications ---------------------------------------------------
    certs = cv.get("certifications") or []
    if certs:
        _heading(doc, headings["certifications"])
        for cert in certs:
            doc.add_paragraph(str(cert).strip(), style="List Bullet")

    # Document title metadata (filename still carries the real name downstream).
    try:
        doc.core_properties.title = str(cv.get("name", "")).strip()
    except Exception:
        pass

    return doc


def main(argv=None) -> int:
    try:  # keep Cyrillic status prints from crashing on a cp1252 console (Windows)
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    parser = argparse.ArgumentParser(description="Render a CV JSON to an ATS-clean .docx")
    parser.add_argument("cv_json", help="Path to the CV JSON file")
    parser.add_argument("--out", required=True, help="Output .docx path")
    args = parser.parse_args(argv)

    cv = json.loads(Path(args.cv_json).read_text(encoding="utf-8"))
    doc = build_document(cv)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"Wrote {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
